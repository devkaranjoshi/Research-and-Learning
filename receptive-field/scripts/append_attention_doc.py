"""Append the ERF / sliding-window experiment as a case-study section to
attention-evolution.docx. Non-destructive: backs up the original first.

Adds: theory, methodology (gradient ruler D), the mathematics of gradient-based ERF,
results (tables + embedded figures), conclusions, limitations.
"""
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCX = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("attention-evolution.docx")
FIG = Path(__file__).resolve().parent.parent / "docs" / "figures"

# --- backup -----------------------------------------------------------------
backup = DOCX.with_suffix(".backup.docx")
if not backup.exists():
    shutil.copy2(DOCX, backup)

doc = Document(str(DOCX))

# Resolve style objects by name once (the doc has a duplicate "Heading 1" name that
# breaks python-docx's name-based lookup, so we assign style OBJECTS directly).
_style_by_name = {}
for _s in doc.styles:
    try:
        if _s.name and _s.name not in _style_by_name:
            _style_by_name[_s.name] = _s
    except Exception:
        pass


# --- helpers ----------------------------------------------------------------
def H(text, level=2):
    p = doc.add_paragraph()
    p.add_run(text)
    st = _style_by_name.get(f"Heading {level}")
    if st is not None:
        p.style = st
    return p


def P(text="", italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p


def MATH(text):
    """Centered equation line (Unicode math)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)
    return p


def BULLET(text):
    p = doc.add_paragraph()
    p.add_run(text)
    st = _style_by_name.get("List Paragraph")
    if st is not None:
        p.style = st
    return p


def _set_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "999999")
        b.append(e)
    tblPr.append(b)


def TABLE(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    _set_borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9)
    doc.add_paragraph()


def FIGURE(name, caption):
    path = FIG / name
    if not path.exists():
        P(f"[missing figure: {name}]", italic=True)
        return
    doc.add_picture(str(path), width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


# ============================================================================
doc.add_page_break()
H("Case Study — Measuring the Effective Receptive Field of Sliding-Window Attention", level=1)
P("This section reports a hands-on experiment that measures, rather than asserts, how much "
  "context a sliding-window attention model actually uses. It extends the “Sparse and windowed "
  "attention” section above with a concrete, reproducible study on Mistral-7B-v0.1 (the canonical "
  "model with sliding_window = 4096). The throughline: a window of W tokens per layer composes "
  "across depth into a theoretical reach of L×W, but the effective receptive field — the input a "
  "token actually depends on — is a small, exponentially-weighted neighborhood far below that "
  "bound. The gap, which widens with depth, is the dilution of sliding-window attention.")

# ---- 1. Question -----------------------------------------------------------
H("1. The question", level=2)
P("For any output token, how many tokens of actual lookback does it have, and how does "
  "sliding-window attention dilute that influence as depth grows? This is distinct from asking "
  "which tokens were attended to, which are reachable, or which contributed to the final "
  "prediction. We ask a causal-sensitivity question: if we slightly perturb one input token, how "
  "much does a chosen internal representation change?")

# ---- 2. Theory -------------------------------------------------------------
H("2. Theory", level=2)

H("2.1 Theoretical vs effective receptive field", level=3)
P("The theoretical receptive field (TRF) is what the architecture allows — which input positions "
  "can structurally reach a given unit. The effective receptive field (ERF) is what the model "
  "actually uses — the magnitude with which each input influences the unit. In CNNs (Luo et al., "
  "2016) the ERF is a small central Gaussian fraction of the TRF and grows only ~√depth. We test "
  "whether the analogous gap holds for a real, trained sliding-window transformer.")

H("2.2 Sliding-window attention and depth", level=3)
P("In Mistral, a query at position q may attend to a key at position kv if and only if both "
  "conditions hold:")
MATH("kv ≤ q    (causal)        and        kv > q − W    (window, W = 4096)")
P("So each token attends to a rolling band of the last W tokens, giving O(N·W) attention cost "
  "instead of O(N²). One layer reaches back W tokens; stacking L layers composes these windows, "
  "because each layer pulls in what the layer below already gathered:")
MATH("theoretical reach(L) ≈ L × W      (the convolution recurrence r_L = r_{L−1} + (W − 1))")
P("For Mistral, 32 × 4096 = 131,072 ≈ the advertised 128k context. Depth is what turns a 4096 "
  "window into a 131k reach. This reach is exact and is verified independently as graph "
  "reachability (boolean powers of the attention mask): reach grows by exactly W − 1 per layer.")
FIGURE("swa_reach_vs_depth.png",
       "Theoretical receptive field grows linearly with depth: reach back = (W−1)·L, exact to the "
       "token (computed as boolean matrix powers of the sliding-window mask). This is the upper "
       "bound the effective field is measured against.")

H("2.3 The tension: linear reach vs dilution", level=3)
P("Two forces oppose each other as depth grows. Theoretical reach grows linearly (L×W). But to "
  "influence a token d positions away, signal must hop through d/W windows, and every hop averages "
  "its contribution over ~W tokens — so a single far token’s influence is divided again at each "
  "hop, decaying roughly exponentially with distance. The receptive field widens with depth, but "
  "the influence at its edge fades. Measuring that decay is the goal.")

# ---- 3. Methodology --------------------------------------------------------
H("3. Methodology — gradient-based ERF", level=2)
P("We operationalize influence as a gradient (the “ruler D” method). For a chosen output unit y, "
  "the influence of input token j is the magnitude of ∂y/∂e_j, where e_j is token j’s embedding. "
  "Large gradient ⇒ that token shapes the unit; exactly zero ⇒ it is outside the effective field. "
  "Unlike raw attention weights, the gradient accounts for the entire computation path "
  "(attention × values, MLPs, residual stream, RoPE, every layer).")
P("Procedure for one measurement (layer L, target token t, window W):")
BULLET("1. Tokenize a passage to length N, with N > W so the window actually engages "
       "(below the window, SWA equals full attention).")
BULLET("2. Embed and set requires_grad on the embeddings — token ids are integers and are not "
       "differentiable, so the gradient must be taken with respect to the continuous embedding vectors.")
BULLET("3. Forward pass with the sliding-window mask active (SDPA kernel, bf16, all weights frozen).")
BULLET("4. Select the output unit: the residual-stream vector h_{L,t} = hidden_states[L+1][t].")
BULLET("5. Reduce to a scalar y = ‖h_{L,t}‖₂ and backpropagate y.backward().")
BULLET("6. Read influence(j) = ‖ emb.grad[j] ‖₂, the L2 norm over d_model of each token’s "
       "embedding gradient. Masked/future positions receive exactly zero gradient.")
P("Two design choices define the measurement: (C1) we backpropagate from the residual-stream norm "
  "at (L, t) — an “into-this-layer” receptive field, the closest analog to the CNN single-center-"
  "unit probe; (C2) we reduce the per-token gradient by its L2 norm over the 4096 feature "
  "dimensions. Compute choices (bf16, SDPA, frozen weights, optional gradient checkpointing) change "
  "only speed and memory, never which tokens influence the unit or by how much.")

# ---- 4. The mathematics ----------------------------------------------------
H("4. The mathematics of gradient-based ERF", level=2)

H("4.1 The transformer as a function", level=3)
P("With an input sequence of N tokens, embedding gives")
MATH("E = [e₀, e₁, …, e_{N−1}],     e_j ∈ ℝ^{d_model},     d_model = 4096")
P("and the transformer computes the residual-stream representations")
MATH("H = f(E) = [h₀, h₁, …, h_{N−1}],     h_i ∈ ℝ^{4096}")

H("4.2 Target representation and the need for a scalar", level=3)
P("We study layer L and target token t, so the representation of interest is h_{L,t} ∈ ℝ^{4096} — "
  "everything the model has computed about token t up to layer L. A gradient requires a scalar "
  "objective (PyTorch’s .backward() expects a scalar), but h_{L,t} is a vector. We therefore "
  "collapse it to a scalar with the L2 norm:")
MATH("y = ‖h_{L,t}‖₂ = √( Σ_i h_i² )")
P("The norm is chosen because it produces one scalar, depends on every feature dimension, treats "
  "all neurons symmetrically, measures overall representation magnitude, and avoids picking an "
  "arbitrary single neuron.")

H("4.3 Differentiating with respect to the embeddings", level=3)
P("We then compute ∂y/∂E. Because E is a list of per-token vectors, the gradient decomposes per "
  "position:")
MATH("∂y/∂E = [ ∂y/∂e₀ , ∂y/∂e₁ , … , ∂y/∂e_{N−1} ]")
P("We differentiate with respect to the embeddings e_j (continuous vectors in ℝ^{4096}) rather "
  "than the token ids, which are discrete integers and carry no derivative.")

H("4.4 What emb.grad contains", level=3)
P("For a sequence of N tokens, emb has shape [N, 4096]. After y.backward(), PyTorch populates "
  "emb.grad with the same shape [N, 4096]. The slice emb.grad[j] is")
MATH("emb.grad[j] = ∂y / ∂e_j     ∈ ℝ^{4096}")
P("a 4096-dimensional vector whose every component answers: if this embedding dimension is "
  "perturbed slightly, how much does y change? To obtain one influence score per token, we take "
  "the L2 norm across the feature dimension:")
MATH("ERF(j) = ‖ ∂y/∂e_j ‖₂ = ‖ emb.grad[j] ‖₂        (emb.grad.norm(dim=−1))")

H("4.5 Why this measures influence", level=3)
P("If ‖∂y/∂e_3‖ = 1.0 while ‖∂y/∂e_1‖ = 0.05, then a tiny perturbation of token 3 produces a "
  "large change in the target representation while token 1 barely moves it — token 3 is far more "
  "influential. This is a direct statement of first-order sensitivity.")

H("4.6 The chain rule accounts for everything", level=3)
P("Backpropagation applies the chain rule through the whole network:")
MATH("∂y/∂e_j = (∂y/∂h) · (∂h/∂e_j)")
P("so the single gradient automatically incorporates attention scores, value vectors, residual "
  "connections, normalization, MLP blocks, rotary position embeddings, and every intermediate "
  "layer. Nothing is measured separately. This is why gradients are a stronger probe than "
  "attention weights: attention answers “where did the query look?”, gradients answer “what "
  "actually changes the computation?”. A token can receive 40% attention yet have zero gradient "
  "(looked at but unused), or 5% attention yet a large gradient (a little attention carrying "
  "decisive information). Attention ≠ influence.")

H("4.7 The effective receptive field, formally", level=3)
P("Putting it together, the effective receptive field of the target unit with respect to input "
  "token j is")
MATH("ERF(j) = ‖ ∂ ‖h_{L,t}‖₂ / ∂e_j ‖₂")
P("Large values indicate meaningful influence; exact zeros indicate no computational path; small "
  "values indicate reachable-but-weakly-influential positions. This is a causal measure of "
  "influence, not a visualization of attention.")

H("4.8 Metrics derived from the influence vector", level=3)
P("Let p_j = ERF(j) / Σ_k ERF(k) be the normalized influence distribution and d_j = t − j the "
  "distance back. We summarize each measurement with:")
MATH("mean effective lookback = Σ_j p_j · d_j")
MATH("90%-mass radius = min{ D : Σ_{d_j ≤ D} p_j ≥ 0.90 }")
MATH("threshold reach = max{ d_j : ERF(j) > 0.01 · max_k ERF(k) }")
MATH("participation ratio = 1 / Σ_j p_j²     (effective number of tokens attended)")
MATH("dilution ratio = threshold reach / min( (L+1)·W , t )")
P("Multiple metrics are used because threshold reach saturates (a single far spike pins it to N); "
  "the mass-weighted metrics capture the distribution, which is where dilution lives.")

# ---- 5. Results ------------------------------------------------------------
H("5. Results", level=2)
P("Model: Mistral-7B-v0.1 (32 layers, 32 heads / 8 KV via GQA, d_model 4096, W 4096). Hardware: "
  "NVIDIA GH200 96 GB. Probe: gradient ruler D, last-token target, layers swept as the depth dial.")

H("5.1 Information propagation in a real forward pass (mechanism)", level=3)
P("Perturbing a single token’s embedding and measuring how far the change has spread after each "
  "layer makes the window-hopping visible. With the window shrunk to 16 for visibility, the "
  "perturbation front advances by exactly W − 1 = 15 positions per layer, matching the theoretical "
  "front s + L·(W−1) to the token, until it saturates at the sequence end.")
FIGURE("mistral_propagation.png",
       "Real Mistral-7B forward pass (window shrunk to 16 for visibility). A perturbation at the "
       "source token spreads forward by exactly W−1 positions per layer — a sharp staircase "
       "wavefront. Black = exactly zero change (not yet reached). This is the receptive field being "
       "built by depth.")

H("5.2 Native window (W = 4096): effective field ≪ reachable", level=3)
P("With native W = 4096 and a 8000-token prompt, influence reaches the whole prompt by layer 8, "
  "yet the effective field is only ~500 tokens (participation ratio), centered ~3.3k back. "
  "Reachable is not the same as used. The right-most column applies the attention-sink control "
  "(see 5.4).")
TABLE(["layer", "mean lookback", "90%-mass radius", "eff #tokens", "reach (1%)", "reach (sink-masked)"],
      [[0, 490, 2099, 5.7, 13, 13],
       [4, 1401, 4039, 235, 7999, 1293],
       [8, 1919, 4848, 393, 7999, 4083],
       [16, 2789, 5719, 513, 7998, 7998],
       [24, 3073, 6008, 467, 7998, 7998],
       [31, 3376, 6305, 500, 7998, 7998]])
FIGURE("mistral_erf_dilution_W4096.png",
       "Influence vs distance per layer, native W=4096. Layer 0 has a hard cliff at exactly 4096 "
       "(single-layer reach), then the structural-zero floor. Deeper layers leak past 4096 but with "
       "a visible step-down at the window boundary (the extra-hop penalty). The spike at distance "
       "≈8000 is the position-0 attention sink.")

H("5.3 Small window (W = 512): the depth-dilution", level=3)
P("Native W caps the theoretical reach at the sequence length by layer 1, hiding the depth story. "
  "With W = 512, the theoretical reach (L+1)·W grows ~512 per layer for ~8 layers, exposing the "
  "divergence. The dilution ratio (effective reach / theoretical reach) falls from 0.93 at layer 0 "
  "to 0.33 at layer 8: at the first layer the model uses ~93% of its single-window reach, but by "
  "layer 8 the theoretical reach has grown 8× while the effective reach uses only ~33% of it. "
  "Effective lookback grows sublinearly while theoretical reach grows linearly.")
TABLE(["layer", "theoretical (L+1)·W", "effective reach (1%)", "dilution ratio", "mean lookback", "eff #tokens"],
      [[0, 512, 475, 0.93, 127, 16.5],
       [1, 1024, 510, 0.50, 147, 24.7],
       [2, 1536, 655, 0.43, 199, 73.5],
       [4, 2560, 1073, 0.42, 278, 175],
       [8, 3999, 1315, 0.33, 334, 245],
       [16, "3999*", 1755, "—", 534, 547],
       [31, "3999*", 2147, "—", 653, 949]])
P("*theoretical reach capped at the sequence length (4000) from layer 8 onward.", italic=True, size=9)
FIGURE("mistral_erf_dilution_W512.png",
       "Influence vs distance per layer, W=512. Each layer shows a hard cliff at exactly (L+1)·W "
       "(theoretical reach) plus an approximately exponential decay within it (the dilution). Deeper "
       "layers push the cliff outward and soften the slope, but the far tail is heavily diluted.")

H("5.4 Attention-sink control", level=3)
P("The first token acts as an attention sink. Re-running with its influence zeroed shows the "
  "mass-weighted metrics (mean lookback, 90%-mass radius, participation ratio) move by less than "
  "3% — the dilution finding is robust. However, the threshold-reach metric was sink-inflated at "
  "shallow layers (layer 4: 7999 → 1293; layer 8: 7999 → 4083). Removing the sink reveals the "
  "genuine effective reach growing with depth (1293 → 4083 → 7998). The control kills a metric, not "
  "the finding: the gradient is causally correct (the sink truly influences), but its single far "
  "spike inflates max-relative metrics while barely touching mass-weighted ones.")

# ---- 6. Conclusions --------------------------------------------------------
H("6. Conclusions", level=2)
P("A token’s effective receptive field in Mistral-7B is a small, exponentially-weighted "
  "neighborhood — about 500 tokens (mean lookback ≈ 3.3k) under the native 4096 window — far below "
  "the 8000 reachable in the prompt and the 131,072 theoretically reachable. Theoretical reach "
  "grows linearly with depth ((L+1)·W); effective lookback grows sublinearly; the widening gap is "
  "the dilution of sliding-window attention, caused by each window-hop dividing a far token’s "
  "influence across ~W tokens. This reproduces the convolutional effective-receptive-field law "
  "(effective ≪ theoretical) for a real, trained sliding-window transformer, and makes precise the "
  "sense in which windowed attention “approximates” full attention: it reconstructs the local "
  "effective field cheaply, while the long tail it nominally reaches is heavily diluted.")

# ---- 7. Limitations --------------------------------------------------------
H("7. Limitations", level=2)
P("This is the gradient (first-order) notion of influence — not attention weights and not a full "
  "nonlinear ablation, though it is the most faithful cheap proxy for causal influence. Results are "
  "from a single prompt with a last-token target (not a corpus statistic); generalization would "
  "average over prompts and targets. The bf16 noise floor is ~1e-3; structural zeros beyond the "
  "mask are exact, but tiny in-field values near the floor are not precise. Threshold reach "
  "saturates and is sink-sensitive — the mass-weighted metrics are the trustworthy ones.")

# --- save -------------------------------------------------------------------
try:
    doc.save(str(DOCX))
    print(f"SAVED -> {DOCX}")
except PermissionError:
    alt = DOCX.with_name("attention-evolution-with-experiment.docx")
    doc.save(str(alt))
    print(f"ORIGINAL LOCKED (open in Word?). SAVED -> {alt}")
print(f"backup at -> {backup}")
